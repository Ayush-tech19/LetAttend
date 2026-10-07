import os
import io
import json
import sqlite3
import hashlib
from datetime import date, datetime

import cv2
import numpy as np
import streamlit as st
from PIL import Image, ImageOps
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment


# ============================================================
# PATHS
# ============================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
MODEL_DIR = os.path.join(BASE_DIR, "model")
DATA_DIR = os.path.join(BASE_DIR, "data")
ATTENDANCE_DIR = os.path.join(BASE_DIR, "Attendance")

YUNET_PATH = os.path.join(MODEL_DIR, "yunet.onnx")
SFACE_PATH = os.path.join(MODEL_DIR, "sface.onnx")
DB_PATH = os.path.join(DATA_DIR, "attendance.db")
EXCEL_PATH = os.path.join(ATTENDANCE_DIR, "Attendance.xlsx")

os.makedirs(MODEL_DIR, exist_ok=True)
os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(ATTENDANCE_DIR, exist_ok=True)


# ============================================================
# SETTINGS
# ============================================================

DETECTION_THRESHOLD = 0.45
NMS_THRESHOLD = 0.30
TOP_K = 5000

# SFace cosine-similarity acceptance boundary.
# This is an experimentally selected project setting, not a
# universal SFace constant.
RECOGNITION_THRESHOLD = 0.45
MIN_SCORE_MARGIN = 0.03

# The requested registration flow uses exactly five camera photos.
REGISTRATION_PHOTOS = 5
MAX_IMAGE_SIDE = 1600
MIN_REGISTRATION_FACE_SIZE = 55


# ============================================================
# PAGE
# ============================================================

st.set_page_config(
    page_title="Smart Attendance System",
    page_icon="🎓",
    layout="wide"
)


# ============================================================
# SQLITE DATABASE
# ============================================================

def db_connection():
    conn = sqlite3.connect(
        DB_PATH,
        timeout=30,
        check_same_thread=False
    )
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA journal_mode = WAL")
    return conn


def initialize_database():
    conn = db_connection()
    try:
        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS students (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                college_id TEXT NOT NULL UNIQUE,
                name TEXT NOT NULL,
                year TEXT NOT NULL,
                branch TEXT DEFAULT '',
                section TEXT DEFAULT '',
                created_at TEXT NOT NULL
            )
            """
        )

        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS face_embeddings (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                embedding TEXT NOT NULL,
                FOREIGN KEY(student_id)
                    REFERENCES students(id)
                    ON DELETE CASCADE
            )
            """
        )

        conn.execute(
            """
            CREATE TABLE IF NOT EXISTS attendance (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                student_id INTEGER NOT NULL,
                attendance_date TEXT NOT NULL,
                status TEXT NOT NULL DEFAULT 'PRESENT',
                similarity REAL NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL,
                UNIQUE(student_id, attendance_date),
                FOREIGN KEY(student_id)
                    REFERENCES students(id)
                    ON DELETE CASCADE
            )
            """
        )
        conn.commit()
    finally:
        conn.close()


initialize_database()


# ============================================================
# DATABASE FUNCTIONS
# ============================================================

def get_students():
    conn = db_connection()
    try:
        rows = conn.execute(
            """
            SELECT id, college_id, name, year, branch, section, created_at
            FROM students
            ORDER BY id
            """
        ).fetchall()

        return [
            {
                "id": row[0],
                "college_id": row[1],
                "name": row[2],
                "year": row[3],
                "branch": row[4] or "",
                "section": row[5] or "",
                "created_at": row[6]
            }
            for row in rows
        ]
    finally:
        conn.close()


def get_student(student_id):
    conn = db_connection()
    try:
        row = conn.execute(
            """
            SELECT id, college_id, name, year, branch, section, created_at
            FROM students
            WHERE id = ?
            """,
            (int(student_id),)
        ).fetchone()

        if row is None:
            return None

        return {
            "id": row[0],
            "college_id": row[1],
            "name": row[2],
            "year": row[3],
            "branch": row[4] or "",
            "section": row[5] or "",
            "created_at": row[6]
        }
    finally:
        conn.close()


def college_id_exists(college_id):
    conn = db_connection()
    try:
        row = conn.execute(
            "SELECT 1 FROM students WHERE lower(college_id)=lower(?) LIMIT 1",
            (college_id,)
        ).fetchone()
        return row is not None
    finally:
        conn.close()


def get_all_embeddings():
    conn = db_connection()
    try:
        rows = conn.execute(
            """
            SELECT student_id, embedding
            FROM face_embeddings
            ORDER BY student_id, id
            """
        ).fetchall()
    finally:
        conn.close()

    result = {}
    for student_id, embedding_text in rows:
        try:
            result.setdefault(int(student_id), []).append(
                json.loads(embedding_text)
            )
        except (TypeError, ValueError, json.JSONDecodeError):
            continue
    return result


def save_student(college_id, name, year, branch, section, embeddings):
    conn = db_connection()
    try:
        conn.execute("BEGIN")

        cursor = conn.execute(
            """
            INSERT INTO students
            (college_id, name, year, branch, section, created_at)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                college_id,
                name,
                year,
                branch,
                section,
                datetime.now().isoformat(timespec="seconds")
            )
        )

        student_id = int(cursor.lastrowid)

        for embedding in embeddings:
            conn.execute(
                """
                INSERT INTO face_embeddings(student_id, embedding)
                VALUES (?, ?)
                """,
                (student_id, json.dumps(embedding))
            )

        conn.commit()
        return student_id
    except Exception:
        conn.rollback()
        raise
    finally:
        conn.close()


def mark_present(student_id, attendance_date, similarity):
    conn = db_connection()
    try:
        now = datetime.now().isoformat(timespec="seconds")

        cursor = conn.execute(
            """
            INSERT OR IGNORE INTO attendance
            (student_id, attendance_date, status, similarity, created_at)
            VALUES (?, ?, 'PRESENT', ?, ?)
            """,
            (
                int(student_id),
                attendance_date,
                float(similarity),
                now
            )
        )

        if cursor.rowcount == 1:
            conn.commit()
            return "NEW"

        # Keep the strongest score if the same student is detected again.
        conn.execute(
            """
            UPDATE attendance
            SET similarity = CASE
                WHEN similarity < ? THEN ?
                ELSE similarity
            END
            WHERE student_id = ? AND attendance_date = ?
            """,
            (
                float(similarity),
                float(similarity),
                int(student_id),
                attendance_date
            )
        )
        conn.commit()
        return "ALREADY"
    finally:
        conn.close()


def get_today_attendance():
    today = date.today().isoformat()
    conn = db_connection()
    try:
        rows = conn.execute(
            """
            SELECT student_id, status, similarity
            FROM attendance
            WHERE attendance_date = ?
            """,
            (today,)
        ).fetchall()
        return {
            int(row[0]): {
                "status": row[1],
                "similarity": row[2]
            }
            for row in rows
        }
    finally:
        conn.close()


def get_attendance_records():
    conn = db_connection()
    try:
        return conn.execute(
            """
            SELECT
                s.college_id,
                s.name,
                s.year,
                s.branch,
                s.section,
                a.attendance_date,
                a.status,
                a.similarity
            FROM attendance a
            INNER JOIN students s ON s.id = a.student_id
            ORDER BY a.attendance_date DESC, s.name
            """
        ).fetchall()
    finally:
        conn.close()


# ============================================================
# EXCEL REBUILD
# ============================================================

def rebuild_excel():
    """Create Excel from SQLite so Excel cannot become a second database."""
    students = get_students()
    records = get_attendance_records()

    dates = sorted({row[5] for row in records})
    attendance_map = {
        (row[0], row[5]): row[6]
        for row in records
    }

    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"

    headers = [
        "College ID",
        "Name",
        "Year",
        "Branch",
        "Section"
    ] + dates

    ws.append(headers)

    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    for student in students:
        row = [
            student["college_id"],
            student["name"],
            student["year"],
            student["branch"],
            student["section"]
        ]

        for attendance_date in dates:
            row.append(
                attendance_map.get(
                    (student["college_id"], attendance_date),
                    ""
                )
            )

        ws.append(row)

    fixed_widths = {
        1: 16,
        2: 25,
        3: 14,
        4: 18,
        5: 14
    }

    for column_index, width in fixed_widths.items():
        ws.column_dimensions[
            ws.cell(row=1, column=column_index).column_letter
        ].width = width

    for column_index in range(6, ws.max_column + 1):
        ws.column_dimensions[
            ws.cell(row=1, column=column_index).column_letter
        ].width = 14

    try:
        wb.save(EXCEL_PATH)
        return True, None
    except PermissionError:
        return False, (
            "Attendance.xlsx is open. Close Excel and try again. "
            "The database has still been updated safely."
        )


# ============================================================
# IMAGE FUNCTIONS
# ============================================================

def prepare_image(file_bytes):
    """Read mobile images with EXIF rotation handled correctly."""
    try:
        pil_image = Image.open(io.BytesIO(file_bytes))
        pil_image = ImageOps.exif_transpose(pil_image)
        pil_image = pil_image.convert("RGB")

        rgb = np.asarray(pil_image)
        image = cv2.cvtColor(rgb, cv2.COLOR_RGB2BGR)

        height, width = image.shape[:2]
        largest_side = max(height, width)

        if largest_side > MAX_IMAGE_SIDE:
            scale = MAX_IMAGE_SIDE / largest_side
            new_width = max(1, int(width * scale))
            new_height = max(1, int(height * scale))
            image = cv2.resize(
                image,
                (new_width, new_height),
                interpolation=cv2.INTER_AREA
            )

        return image
    except Exception:
        return None


def detect_faces(image):
    height, width = image.shape[:2]
    detector.setInputSize((width, height))
    _, faces = detector.detect(image)
    return [] if faces is None else list(faces)


def create_embedding(image):
    """Registration requires exactly one usable face in the photo."""
    faces = detect_faces(image)

    if len(faces) == 0:
        return None, (
            "No face detected. Move closer, keep your full face visible, "
            "and use good lighting."
        )

    if len(faces) > 1:
        return None, (
            "More than one face detected. Only one person is allowed "
            "in each registration photo."
        )

    face = faces[0]
    face_width = float(face[2])
    face_height = float(face[3])

    if (
        face_width < MIN_REGISTRATION_FACE_SIZE
        or face_height < MIN_REGISTRATION_FACE_SIZE
    ):
        return None, (
            "Face is too small. Move closer to the camera and retake the photo."
        )

    try:
        aligned = recognizer.alignCrop(image, face)
        feature = recognizer.feature(aligned).astype(np.float32)
        return feature.flatten().tolist(), None
    except Exception as e:
        return None, f"Face feature extraction failed: {e}"


# ============================================================
# RECOGNITION
# ============================================================

def normalize_embedding(value):
    arr = np.asarray(value, dtype=np.float32)

    if arr.ndim == 1:
        return [arr.reshape(1, -1)]

    if arr.ndim == 2:
        return [row.reshape(1, -1) for row in arr]

    return []


def recognize_face(image, face, stored_embeddings):
    try:
        aligned = recognizer.alignCrop(image, face)
        query_feature = recognizer.feature(aligned)
    except Exception:
        return None, 0.0, -1.0

    student_scores = []

    for student_id, stored in stored_embeddings.items():
        references = normalize_embedding(stored)
        if not references:
            continue

        best_for_student = -1.0

        for reference in references:
            try:
                score = float(
                    recognizer.match(
                        query_feature,
                        reference,
                        cv2.FaceRecognizerSF_FR_COSINE
                    )
                )
                best_for_student = max(
                    best_for_student,
                    score
                )
            except Exception:
                continue

        if best_for_student >= 0:
            student_scores.append(
                (int(student_id), best_for_student)
            )

    if not student_scores:
        return None, 0.0, -1.0

    student_scores.sort(
        key=lambda item: item[1],
        reverse=True
    )

    best_student, best_score = student_scores[0]
    second_score = (
        student_scores[1][1]
        if len(student_scores) > 1
        else -1.0
    )

    margin_ok = (
        second_score < 0
        or best_score - second_score >= MIN_SCORE_MARGIN
    )

    if best_score >= RECOGNITION_THRESHOLD and margin_ok:
        return best_student, best_score, second_score

    return None, best_score, second_score


# ============================================================
# DRAW BOXES
# ============================================================

def draw_face_label(image, face, label, score=None):
    x, y, w, h = [int(float(v)) for v in face[:4]]

    known = label.strip().lower() != "unknown"
    box_color = (0, 200, 0) if known else (0, 0, 255)

    cv2.rectangle(
        image,
        (x, y),
        (x + w, y + h),
        box_color,
        3
    )

    text = label
    if known and score is not None:
        text = f"{label} ({score:.2f})"

    font = cv2.FONT_HERSHEY_SIMPLEX
    font_scale = 0.65
    thickness = 2
    (tw, th), baseline = cv2.getTextSize(
        text,
        font,
        font_scale,
        thickness
    )

    label_y = max(
        y,
        th + baseline + 10
    )

    cv2.rectangle(
        image,
        (x, label_y - th - baseline - 8),
        (x + tw + 10, label_y + 4),
        box_color,
        -1
    )

    cv2.putText(
        image,
        text,
        (x + 5, label_y - 3),
        font,
        font_scale,
        (255, 255, 255),
        thickness,
        cv2.LINE_AA
    )


# ============================================================
# MODEL LOADING
# ============================================================

@st.cache_resource
def load_models():
    if not os.path.exists(YUNET_PATH):
        raise FileNotFoundError(
            f"YuNet model not found: {YUNET_PATH}"
        )

    if not os.path.exists(SFACE_PATH):
        raise FileNotFoundError(
            f"SFace model not found: {SFACE_PATH}"
        )

    detector = cv2.FaceDetectorYN.create(
        YUNET_PATH,
        "",
        (320, 320),
        DETECTION_THRESHOLD,
        NMS_THRESHOLD,
        TOP_K
    )

    recognizer = cv2.FaceRecognizerSF.create(
        SFACE_PATH,
        ""
    )

    return detector, recognizer


try:
    detector, recognizer = load_models()
except Exception as e:
    st.error("The YuNet/SFace models could not be loaded.")
    st.code(str(e))
    st.stop()


# ============================================================
# APP DATA
# ============================================================

students = get_students()
stored_embeddings = get_all_embeddings()


# ============================================================
# SIDEBAR
# ============================================================

st.sidebar.title("🎓 Smart Attendance")

page = st.sidebar.radio(
    "Menu",
    [
        "Dashboard",
        "Student Registration",
        "Attendance",
        "Registered Students",
        "Attendance Records"
    ]
)

st.sidebar.divider()
st.sidebar.caption(
    "Registration and attendance data are stored in SQLite, "
    "so Streamlit page refreshes do not clear them."
)


# ============================================================
# DASHBOARD
# ============================================================

if page == "Dashboard":
    st.title("🎓 Smart Attendance System")

    today_display = date.today().strftime("%d-%m-%Y")
    attendance_today = get_today_attendance()

    c1, c2, c3 = st.columns(3)
    c1.metric("Registered Students", len(students))
    c2.metric("Present Today", len(attendance_today))
    c3.metric("Date", today_display)

    st.divider()

    st.info(
        "Register each student once using five face photos. "
        "Then capture or upload one class group photo in Attendance. "
        "Recognized students are automatically marked PRESENT for today's date."
    )

    if students:
        rows = []
        for student in students:
            record = attendance_today.get(student["id"])
            rows.append(
                {
                    "College ID": student["college_id"],
                    "Name": student["name"],
                    "Year": student["year"],
                    "Status": "PRESENT" if record else "—"
                }
            )

        st.subheader("Today's Attendance")
        st.dataframe(
            rows,
            use_container_width=True,
            hide_index=True
        )


# ============================================================
# STUDENT REGISTRATION
# ============================================================

elif page == "Student Registration":
    st.title("👨‍🎓 Student Registration")

    st.write(
        "Enter the student details and capture exactly five photos "
        "using the device camera."
    )

    st.info(
        "Recommended order: Front → Slight Left → Slight Right → "
        "Slight Up → Slight Down. Only one person should appear "
        "in each photo."
    )

    registration_round = st.session_state.get(
        "registration_round",
        0
    )

    name = st.text_input(
        "Student Name",
        key=f"reg_name_{registration_round}"
    )

    college_id = st.text_input(
        "College ID / Roll Number",
        key=f"reg_id_{registration_round}"
    )

    year = st.selectbox(
        "Year",
        [
            "1st Year",
            "2nd Year",
            "3rd Year",
            "4th Year"
        ],
        key=f"reg_year_{registration_round}"
    )

    branch = st.text_input(
        "Branch (optional)",
        key=f"reg_branch_{registration_round}"
    )

    section = st.text_input(
        "Section (optional)",
        key=f"reg_section_{registration_round}"
    )

    st.subheader("Capture Five Facial Photos")

    camera_labels = [
        "Photo 1 — Look straight",
        "Photo 2 — Turn slightly left",
        "Photo 3 — Turn slightly right",
        "Photo 4 — Look slightly up",
        "Photo 5 — Look slightly down"
    ]

    camera_photos = []

    for index, label in enumerate(
        camera_labels,
        start=1
    ):
        camera_photos.append(
            st.camera_input(
                label,
                key=(
                    f"reg_camera_"
                    f"{registration_round}_"
                    f"{index}"
                )
            )
        )

    captured_count = sum(
        photo is not None
        for photo in camera_photos
    )

    st.write(
        f"**Photos captured: {captured_count}/5**"
    )

    register_clicked = st.button(
        "✅ Register Student",
        type="primary",
        use_container_width=True
    )

    if register_clicked:
        clean_name = name.strip()
        clean_id = college_id.strip()
        clean_branch = branch.strip()
        clean_section = section.strip()

        if not clean_name:
            st.error("Please enter the student name.")
            st.stop()

        if not clean_id:
            st.error("Please enter the College ID.")
            st.stop()

        if college_id_exists(clean_id):
            st.error(
                "This College ID is already registered."
            )
            st.stop()

        if captured_count != REGISTRATION_PHOTOS:
            st.error(
                "Exactly five facial photos are required."
            )
            st.stop()

        embeddings_for_student = []
        registration_ok = True

        progress = st.progress(0)
        status = st.empty()

        for index, photo in enumerate(
            camera_photos,
            start=1
        ):
            status.info(
                f"Processing registration photo {index}/5..."
            )

            image = prepare_image(
                photo.getvalue()
            )

            if image is None:
                st.error(
                    f"Photo {index} could not be read. Please retake it."
                )
                registration_ok = False
                break

            embedding, error = create_embedding(
                image
            )

            if embedding is None:
                st.error(
                    f"Photo {index}: {error}"
                )
                registration_ok = False
                break

            embeddings_for_student.append(
                embedding
            )

            progress.progress(
                index / REGISTRATION_PHOTOS
            )

        status.empty()
        progress.empty()

        # Nothing is written until all five photos pass validation.
        if not registration_ok:
            st.error(
                "Registration cancelled. No partial student record was saved."
            )
            st.stop()

        try:
            save_student(
                clean_id,
                clean_name,
                year,
                clean_branch,
                clean_section,
                embeddings_for_student
            )

            excel_ok, excel_message = rebuild_excel()

            st.success(
                f"{clean_name} registered successfully."
            )
            st.info(
                "Five SFace face embeddings have been saved. "
                "Registration photos are not stored."
            )

            if not excel_ok:
                st.warning(excel_message)

            # Change widget keys so the next registration starts fresh.
            st.session_state.registration_round = (
                registration_round + 1
            )
            st.rerun()

        except sqlite3.IntegrityError:
            st.error(
                "This College ID is already registered."
            )
        except Exception as e:
            st.error(
                "Registration could not be completed."
            )
            st.code(str(e))


# ============================================================
# ATTENDANCE
# ============================================================

elif page == "Attendance":
    st.title("📸 Automatic Group Attendance")

    if not students:
        st.warning(
            "No students are registered yet. Register students first."
        )
    else:
        st.write(
            "Upload or capture one class/group photo. The system will "
            "detect every face, identify registered students, draw a box "
            "with their name, label non-registered faces as Unknown, and "
            "automatically mark every recognized student PRESENT for today's date."
        )

        uploaded = st.file_uploader(
            "Upload Group Photo",
            type=["jpg", "jpeg", "png"],
            key="attendance_upload"
        )

        camera = st.camera_input(
            "Or take a Group Attendance Photo",
            key="attendance_camera"
        )

        selected_file = (
            uploaded
            if uploaded is not None
            else camera
        )

        if selected_file is not None:
            file_bytes = selected_file.getvalue()
            image = prepare_image(file_bytes)

            if image is None:
                st.error("Unable to read the image.")
                st.stop()

            faces = detect_faces(image)

            if not faces:
                st.warning(
                    "No faces were detected. Use a clearer group photo "
                    "with faces visible and reasonably well lit."
                )
                st.stop()

            st.info(
                f"{len(faces)} face(s) detected."
            )

            annotated = image.copy()
            result_rows = []
            recognized_students = {}
            recognized_scores = {}

            for face in faces:
                student_id, score, _ = recognize_face(
                    image,
                    face,
                    stored_embeddings
                )

                if student_id is not None:
                    student = get_student(student_id)

                    if student is not None:
                        # Same student may appear more than once; use strongest match.
                        recognized_students[student_id] = student
                        recognized_scores[student_id] = max(
                            score,
                            recognized_scores.get(student_id, -1.0)
                        )

                        draw_face_label(
                            annotated,
                            face,
                            student["name"],
                            score
                        )

                        result_rows.append(
                            {
                                "Student": student["name"],
                                "College ID": student["college_id"],
                                "Similarity": round(score, 3),
                                "Result": "Recognized"
                            }
                        )
                        continue

                draw_face_label(
                    annotated,
                    face,
                    "Unknown"
                )

                result_rows.append(
                    {
                        "Student": "Unknown",
                        "College ID": "",
                        "Similarity": "",
                        "Result": "Not Registered"
                    }
                )

            st.subheader("Face Recognition Result")
            st.image(
                cv2.cvtColor(
                    annotated,
                    cv2.COLOR_BGR2RGB
                ),
                caption=(
                    "Green box = recognized registered student | "
                    "Red box = Unknown"
                ),
                use_container_width=True
            )

            st.dataframe(
                result_rows,
                use_container_width=True,
                hide_index=True
            )

            attendance_date = date.today().isoformat()
            attendance_display = date.today().strftime(
                "%d-%m-%Y"
            )

            if recognized_students:
                # Prevent repeated Streamlit reruns from producing repeated messages.
                photo_token = hashlib.sha256(
                    file_bytes
                ).hexdigest()
                attendance_token = (
                    attendance_date + ":" + photo_token
                )

                already_processed = (
                    st.session_state.get(
                        "last_attendance_token"
                    )
                    == attendance_token
                )

                new_count = 0
                already_count = 0

                if not already_processed:
                    for student_id in recognized_students:
                        result = mark_present(
                            student_id,
                            attendance_date,
                            recognized_scores[student_id]
                        )

                        if result == "NEW":
                            new_count += 1
                        else:
                            already_count += 1

                    st.session_state[
                        "last_attendance_token"
                    ] = attendance_token

                    excel_ok, excel_message = rebuild_excel()

                    if new_count:
                        st.success(
                            f"{new_count} student(s) automatically marked "
                            f"PRESENT for {attendance_display}."
                        )

                    if already_count:
                        st.info(
                            f"{already_count} student(s) were already PRESENT "
                            "today. No duplicate attendance was created."
                        )

                    if not excel_ok:
                        st.warning(excel_message)
                else:
                    st.success(
                        f"Attendance already processed for this photo today. "
                        f"Recognized students are PRESENT for {attendance_display}."
                    )

            else:
                st.warning(
                    "No registered student was confidently recognized. "
                    "No attendance was marked."
                )


# ============================================================
# REGISTERED STUDENTS
# ============================================================

elif page == "Registered Students":
    st.title("👨‍🎓 Registered Students")

    students = get_students()

    if not students:
        st.info("No students are registered yet.")
    else:
        search = st.text_input(
            "Search by name or College ID"
        ).strip().lower()

        filtered = []
        for student in students:
            if (
                not search
                or search in student["name"].lower()
                or search in student["college_id"].lower()
            ):
                filtered.append(student)

        st.write(
            f"Showing {len(filtered)} student(s)."
        )

        for student in filtered:
            with st.container(border=True):
                st.subheader(student["name"])
                c1, c2, c3, c4 = st.columns(4)
                c1.write(
                    "**College ID**\n\n" + student["college_id"]
                )
                c2.write(
                    "**Year**\n\n" + student["year"]
                )
                c3.write(
                    "**Branch**\n\n" + (student["branch"] or "—")
                )
                c4.write(
                    "**Section**\n\n" + (student["section"] or "—")
                )


# ============================================================
# ATTENDANCE RECORDS
# ============================================================

elif page == "Attendance Records":
    st.title("📊 Attendance Records")

    records = get_attendance_records()

    if records:
        table = []
        for row in records:
            table.append(
                {
                    "College ID": row[0],
                    "Name": row[1],
                    "Year": row[2],
                    "Branch": row[3],
                    "Section": row[4],
                    "Date": row[5],
                    "Status": row[6],
                    "Similarity": round(row[7], 3)
                }
            )

        st.dataframe(
            table,
            use_container_width=True,
            hide_index=True
        )
    else:
        st.info("No attendance records are available yet.")

    if os.path.exists(EXCEL_PATH):
        try:
            with open(EXCEL_PATH, "rb") as f:
                st.download_button(
                    "⬇️ Download Attendance Excel",
                    data=f,
                    file_name="Attendance.xlsx",
                    mime=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    )
                )
        except Exception:
            pass

    st.caption(
        "Only recognized students are marked PRESENT. Faces not recognized "
        "in a group photo are left without attendance; they are not marked absent."
    )
