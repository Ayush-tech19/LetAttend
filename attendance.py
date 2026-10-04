import os
import json
import cv2
import numpy as np
import streamlit as st
from datetime import date
from openpyxl import Workbook, load_workbook


# =========================================================
# PATHS
# =========================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))

YUNET_PATH = os.path.join(BASE_DIR, "model", "yunet.onnx")
SFACE_PATH = os.path.join(BASE_DIR, "model", "sface.onnx")
EMBEDDINGS_PATH = os.path.join(BASE_DIR, "model", "embeddings.json")

ATTENDANCE_DIR = os.path.join(BASE_DIR, "Attendance")
EXCEL_PATH = os.path.join(ATTENDANCE_DIR, "Attendance.xlsx")

STUDENT_INFO_PATH = os.path.join(BASE_DIR, "students.json")


# =========================================================
# SETTINGS
# =========================================================

# Increased from 0.40 to reduce false positives
RECOGNITION_THRESHOLD = 0.50


# =========================================================
# PAGE CONFIG
# =========================================================

st.set_page_config(
    page_title="College Attendance System",
    page_icon="🎓",
    layout="wide"
)


# =========================================================
# LOAD MODELS
# =========================================================

@st.cache_resource
def load_models():

    detector = cv2.FaceDetectorYN.create(
        YUNET_PATH,
        "",
        (320, 320),
        0.6,
        0.3,
        5000
    )

    recognizer = cv2.FaceRecognizerSF.create(
        SFACE_PATH,
        ""
    )

    return detector, recognizer


# =========================================================
# LOAD STUDENT DATA
# =========================================================

def load_students():

    if not os.path.exists(STUDENT_INFO_PATH):
        return {}

    with open(STUDENT_INFO_PATH, "r", encoding="utf-8") as f:
        data = json.load(f)

    students = {}

    for student in data:

        student_id = student.get("student_id")

        if student_id:
            students[student_id] = student

    return students


# =========================================================
# LOAD EMBEDDINGS
# =========================================================

def load_embeddings():

    if not os.path.exists(EMBEDDINGS_PATH):
        return {}

    with open(EMBEDDINGS_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


# =========================================================
# COSINE SIMILARITY
# =========================================================

def calculate_similarity(recognizer, feature1, feature2):

    return recognizer.match(
        feature1,
        feature2,
        cv2.FaceRecognizerSF_FR_COSINE
    )


# =========================================================
# RECOGNIZE ONE FACE
# =========================================================

def recognize_face(
    face,
    image,
    detector,
    recognizer,
    embeddings,
    students
):

    aligned_face = recognizer.alignCrop(image, face)

    feature = recognizer.feature(aligned_face)

    best_student = None
    best_score = -1

    for student_id, reference_embedding in embeddings.items():

        try:

            reference = np.array(
                reference_embedding,
                dtype=np.float32
            )

            # Make sure shape is correct
            reference = reference.reshape(1, -1)

            score = calculate_similarity(
                recognizer,
                feature,
                reference
            )

            if score > best_score:

                best_score = score
                best_student = student_id

        except Exception:
            continue

    # Strict threshold
    if best_score >= RECOGNITION_THRESHOLD:

        if best_student in students:

            return best_student, best_score

    return None, best_score


# =========================================================
# CREATE / LOAD EXCEL
# =========================================================

def create_excel_if_needed(students):

    os.makedirs(ATTENDANCE_DIR, exist_ok=True)

    if not os.path.exists(EXCEL_PATH):

        wb = Workbook()
        ws = wb.active
        ws.title = "Attendance"

        headers = [
            "College ID",
            "Name",
            "Branch",
            "Year",
            "Section"
        ]

        ws.append(headers)

        for student in students.values():

            ws.append([
                student.get("college_id", ""),
                student.get("name", ""),
                student.get("branch", ""),
                student.get("year", ""),
                student.get("section", "")
            ])

        wb.save(EXCEL_PATH)


# =========================================================
# SAVE ATTENDANCE
# =========================================================

def save_attendance(present_students, students):

    create_excel_if_needed(students)

    try:

        wb = load_workbook(EXCEL_PATH)

        ws = wb["Attendance"]

        today = date.today().strftime("%d-%m-%Y")

        # Find today's column
        date_column = None

        for col in range(1, ws.max_column + 1):

            if ws.cell(row=1, column=col).value == today:

                date_column = col
                break

        # Create today's column if not present
        if date_column is None:

            date_column = ws.max_column + 1

            ws.cell(
                row=1,
                column=date_column
            ).value = today

        # Create mapping using College ID
        student_rows = {}

        for row in range(2, ws.max_row + 1):

            college_id = ws.cell(
                row=row,
                column=1
            ).value

            if college_id is not None:

                student_rows[str(college_id)] = row

        # Mark ONLY recognized students
        for student_id in present_students:

            student = students.get(student_id)

            if not student:
                continue

            college_id = str(
                student.get("college_id", "")
            )

            if college_id in student_rows:

                row = student_rows[college_id]

                ws.cell(
                    row=row,
                    column=date_column
                ).value = "PRESENT"

        wb.save(EXCEL_PATH)

        return True, None

    except PermissionError:

        return False, (
            "Attendance.xlsx is currently open. "
            "Please close the Excel file and try again."
        )

    except Exception as e:

        return False, str(e)


# =========================================================
# HEADER
# =========================================================

st.title("🎓 College Face Attendance System")

st.caption(
    "Upload a single or group photo to automatically mark "
    "recognized students as present."
)


# =========================================================
# SIDEBAR
# =========================================================

st.sidebar.title("Attendance Menu")

st.sidebar.info(
    "Only students recognized in the uploaded photo "
    "will be marked PRESENT."
)


# =========================================================
# LOAD DATA
# =========================================================

try:

    detector, recognizer = load_models()

except Exception as e:

    st.error("Unable to load face recognition models.")

    st.exception(e)

    st.stop()


students = load_students()
embeddings = load_embeddings()


# =========================================================
# CHECK DATABASE
# =========================================================

if not students:

    st.warning(
        "No student information was found."
    )

    st.stop()


if not embeddings:

    st.warning(
        "No face embeddings were found."
    )

    st.stop()


# =========================================================
# IMAGE INPUT
# =========================================================

st.subheader("Upload Attendance Photo")

uploaded_file = st.file_uploader(
    "Choose a photo",
    type=["jpg", "jpeg", "png"]
)


camera_file = st.camera_input(
    "Or take a photo using camera"
)


image_file = uploaded_file if uploaded_file else camera_file


# =========================================================
# PROCESS PHOTO
# =========================================================

if image_file is not None:

    file_bytes = np.asarray(
        bytearray(image_file.read()),
        dtype=np.uint8
    )

    image = cv2.imdecode(
        file_bytes,
        cv2.IMREAD_COLOR
    )

    if image is None:

        st.error("Unable to read the uploaded image.")

        st.stop()

    # Set detector input size
    height, width = image.shape[:2]

    detector.setInputSize(
        (width, height)
    )

    # Detect faces
    _, detections = detector.detect(image)

    if detections is None:

        st.warning(
            "No faces were detected in the photo."
        )

        st.stop()

    face_count = len(detections)

    st.info(
        f"{face_count} face(s) detected in the photo."
    )

    # =====================================================
    # RECOGNITION
    # =====================================================

    recognized_students = {}

    for face in detections:

        student_id, score = recognize_face(
            face,
            image,
            detector,
            recognizer,
            embeddings,
            students
        )

        if student_id is not None:

            # Avoid duplicate attendance
            if student_id not in recognized_students:

                recognized_students[student_id] = score


    # =====================================================
    # RESULT
    # =====================================================

    st.divider()

    if recognized_students:

        st.success(
            f"{len(recognized_students)} student(s) recognized."
        )

        st.subheader("Present Students")

        for student_id in recognized_students:

            student = students[student_id]

            college_id = student.get(
                "college_id",
                ""
            )

            name = student.get(
                "name",
                "Unknown"
            )

            branch = student.get(
                "branch",
                ""
            )

            year = student.get(
                "year",
                ""
            )

            section = student.get(
                "section",
                ""
            )

            st.markdown(
                f"""
                **{name}**

                College ID: `{college_id}`  
                Branch: `{branch}`  
                Year: `{year}`  
                Section: `{section}`
                """
            )

            st.divider()

    else:

        st.warning(
            "No registered student was confidently recognized."
        )


    # =====================================================
    # MARK ATTENDANCE BUTTON
    # =====================================================

    if st.button(
        "✅ Mark Attendance",
        use_container_width=True
    ):

        if not recognized_students:

            st.warning(
                "No student was recognized. "
                "Attendance was not marked."
            )

        else:

            success, error = save_attendance(
                recognized_students,
                students
            )

            if success:

                st.success(
                    f"Attendance successfully marked for "
                    f"{len(recognized_students)} student(s)."
                )

                st.info(
                    "Students who were not recognized were "
                    "left blank and were NOT marked absent."
                )

            else:

                st.error(
                    f"Attendance could not be saved: {error}"
                )